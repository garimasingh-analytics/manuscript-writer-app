from fastapi import FastAPI, APIRouter, HTTPException, Depends, UploadFile, File, Form
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from fastapi.responses import StreamingResponse
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, EmailStr, model_validator
from typing import List, Optional, Dict, Any
import uuid
from datetime import datetime, timezone, timedelta
import jwt
import bcrypt
import httpx
import asyncio
import io
import json
import re
import anthropic
from PyPDF2 import PdfReader
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# JWT
JWT_SECRET = os.environ.get('JWT_SECRET', 'manuscript-writer-secret-key')
JWT_ALGORITHM = "HS256"
JWT_EXPIRATION_HOURS = 24

# Anthropic client
anthropic_client = anthropic.AsyncAnthropic(api_key=os.environ['ANTHROPIC_API_KEY'])

app = FastAPI(title="Manuscript Writer API")

api_router = APIRouter(prefix="/api")
auth_router = APIRouter(prefix="/auth", tags=["Auth"])
projects_router = APIRouter(prefix="/projects", tags=["Projects"])
agents_router = APIRouter(prefix="/agents", tags=["Agents"])
export_router = APIRouter(prefix="/export", tags=["Export"])
keywords_router = APIRouter(prefix="/keywords", tags=["Keywords"])

security = HTTPBearer()
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# ==================== CLAUDE HELPER ====================

async def call_claude(prompt: str, system: str = "You are a helpful research assistant.", model: str = "claude-sonnet-4-6") -> str:
    message = await anthropic_client.messages.create(
        model=model,
        max_tokens=4096,
        system=system,
        messages=[{"role": "user", "content": prompt}]
    )
    return message.content[0].text

# ==================== DATABASES ====================

AVAILABLE_DATABASES = {
    "pubmed": {"name": "PubMed", "description": "Biomedical literature from MEDLINE", "category": "biomedical", "enabled_by_default": True, "supports_mesh": True},
    "pmc": {"name": "PubMed Central", "description": "Free full-text biomedical archive", "category": "biomedical", "enabled_by_default": True, "supports_mesh": True},
    "europe_pmc": {"name": "Europe PMC", "description": "European life sciences repository", "category": "biomedical", "enabled_by_default": True, "supports_mesh": False},
    "semantic_scholar": {"name": "Semantic Scholar", "description": "AI-powered research database", "category": "multidisciplinary", "enabled_by_default": True, "supports_mesh": False},
    "crossref": {"name": "CrossRef", "description": "DOI registration and scholarly metadata", "category": "multidisciplinary", "enabled_by_default": True, "supports_mesh": False},
    "doaj": {"name": "DOAJ", "description": "Directory of Open Access Journals", "category": "open_access", "enabled_by_default": True, "supports_mesh": False},
    "biorxiv": {"name": "bioRxiv", "description": "Preprint server for biology", "category": "preprint", "enabled_by_default": True, "supports_mesh": False},
    "medrxiv": {"name": "medRxiv", "description": "Preprint server for health sciences", "category": "preprint", "enabled_by_default": True, "supports_mesh": False},
    "arxiv": {"name": "arXiv", "description": "Preprints for physics/math/CS", "category": "preprint", "enabled_by_default": False, "supports_mesh": False},
    "zenodo": {"name": "Zenodo", "description": "Open research repository", "category": "repository", "enabled_by_default": False, "supports_mesh": False},
}

# ==================== MODELS ====================

class UserCreate(BaseModel):
    email: EmailStr
    password: str
    name: str

class UserLogin(BaseModel):
    email: EmailStr
    password: str

class UserResponse(BaseModel):
    id: str
    email: str
    name: str
    created_at: str

class TokenResponse(BaseModel):
    access_token: str
    token_type: str = "bearer"
    user: UserResponse

class StudySummary(BaseModel):
    population_setting: str = ""
    sample_size: str = ""
    study_design: str = ""
    inclusion_criteria: str = ""
    exclusion_criteria: str = ""
    interventions: str = ""
    comparators: str = ""
    outcomes: str = ""
    effect_sizes: str = ""
    follow_up: str = ""
    conclusions: str = ""
    tables: List[Dict[str, Any]] = []
    figures: List[Dict[str, Any]] = []
    charts: List[Dict[str, Any]] = []
    raw_tables_text: str = ""

    @model_validator(mode="before")
    @classmethod
    def _coerce_types(cls, data):
        """Tolerate LLM output that returns null or wrong types for fields."""
        if not isinstance(data, dict):
            return data
        str_fields = {"population_setting", "sample_size", "study_design", "inclusion_criteria",
                      "exclusion_criteria", "interventions", "comparators", "outcomes",
                      "effect_sizes", "follow_up", "conclusions", "raw_tables_text"}
        list_fields = {"tables", "figures", "charts"}
        cleaned = {}
        for k, v in data.items():
            if k in str_fields:
                cleaned[k] = "" if v is None else (v if isinstance(v, str) else str(v))
            elif k in list_fields:
                if isinstance(v, list):
                    cleaned[k] = [item if isinstance(item, dict) else {"value": str(item)} for item in v]
                else:
                    cleaned[k] = []
            else:
                cleaned[k] = v
        return cleaned

class KeywordSet(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    main_keywords: List[str] = []
    exclusion_keywords: List[str] = []
    mesh_terms: List[str] = []
    custom_query: str = ""
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

class SearchFilters(BaseModel):
    main_keywords: List[str] = []
    exclusion_keywords: List[str] = []
    mesh_terms: List[str] = []
    custom_query: str = ""
    year_from: Optional[int] = None
    year_to: Optional[int] = None
    article_types: List[str] = []
    language: str = "english"
    open_access_only: bool = False
    enabled_databases: List[str] = []

class Paper(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    title: str
    authors: str
    year: str
    journal: str
    doi: Optional[str] = None
    pmid: Optional[str] = None
    pmcid: Optional[str] = None
    arxiv_id: Optional[str] = None
    url: Optional[str] = None
    abstract: str = ""
    classification: str = "background"
    article_type: str = "unknown"
    is_open_access: bool = False
    full_text_available: bool = False
    database_source: str = ""
    databases_found_in: List[str] = []
    selected: bool = False
    needs_verification: bool = False
    relevance_score: float = 0.0
    user_notes: str = ""

class ProjectCreate(BaseModel):
    title: str
    description: Optional[str] = ""

class ProjectUpdate(BaseModel):
    title: Optional[str] = None
    description: Optional[str] = None
    study_summary: Optional[StudySummary] = None
    status: Optional[str] = None

class Manuscript(BaseModel):
    title: str = ""
    abstract: str = ""
    introduction: str = ""
    methods: str = ""
    results: str = ""
    evidence_comparison: str = ""
    discussion: str = ""
    conclusion: str = ""
    tables_figures: str = ""
    tables: List[Dict[str, Any]] = []
    figures: List[Dict[str, Any]] = []
    charts: List[Dict[str, Any]] = []
    references: List[Dict[str, Any]] = []
    reference_mapping: Dict[str, List[int]] = {}
    citation_style: str = "endnote"
    version: int = 1
    generated_at: str = ""

class GenerationJob(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    project_id: str
    user_id: str
    status: str = "pending"
    citation_style: str = "endnote"
    sections_to_generate: Optional[List[str]] = None
    word_counts: Optional[Dict[str, int]] = None
    result: Optional[Dict[str, Any]] = None
    error: Optional[str] = None
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())
    completed_at: Optional[str] = None

class ManuscriptGenerateRequest(BaseModel):
    word_counts: Optional[Dict[str, int]] = None

# ==================== AUTH HELPERS ====================

def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode(), bcrypt.gensalt()).decode()

def verify_password(password: str, hashed: str) -> bool:
    return bcrypt.checkpw(password.encode(), hashed.encode())

def create_token(user_id: str) -> str:
    exp = datetime.now(timezone.utc) + timedelta(hours=JWT_EXPIRATION_HOURS)
    return jwt.encode({"sub": user_id, "exp": exp}, JWT_SECRET, algorithm=JWT_ALGORITHM)

async def get_current_user(credentials: HTTPAuthorizationCredentials = Depends(security)):
    try:
        payload = jwt.decode(credentials.credentials, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        user_id = payload.get("sub")
        if not user_id:
            raise HTTPException(status_code=401, detail="Invalid token")
        user = await db.users.find_one({"id": user_id}, {"_id": 0, "password": 0})
        if not user:
            raise HTTPException(status_code=401, detail="User not found")
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid token")

# ==================== AUTH ROUTES ====================

@auth_router.post("/register", response_model=TokenResponse)
async def register(user_data: UserCreate):
    if await db.users.find_one({"email": user_data.email}):
        raise HTTPException(status_code=400, detail="Email already registered")
    user_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    await db.users.insert_one({"id": user_id, "email": user_data.email, "password": hash_password(user_data.password), "name": user_data.name, "created_at": now})
    return TokenResponse(access_token=create_token(user_id), user=UserResponse(id=user_id, email=user_data.email, name=user_data.name, created_at=now))

@auth_router.post("/login", response_model=TokenResponse)
async def login(credentials: UserLogin):
    user = await db.users.find_one({"email": credentials.email})
    if not user or not verify_password(credentials.password, user["password"]):
        raise HTTPException(status_code=401, detail="Invalid credentials")
    return TokenResponse(access_token=create_token(user["id"]), user=UserResponse(id=user["id"], email=user["email"], name=user["name"], created_at=user["created_at"]))

@auth_router.get("/me", response_model=UserResponse)
async def get_me(user: dict = Depends(get_current_user)):
    return UserResponse(**user)

# ==================== PROJECTS ROUTES ====================

@projects_router.get("")
async def list_projects(user: dict = Depends(get_current_user)):
    projects = await db.projects.find({"user_id": user["id"]}, {"_id": 0}).sort("updated_at", -1).to_list(100)
    return projects

@projects_router.post("")
async def create_project(project: ProjectCreate, user: dict = Depends(get_current_user)):
    project_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    doc = {
        "id": project_id,
        "user_id": user["id"],
        "title": project.title,
        "description": project.description,
        "status": "draft",
        "study_summary": None,
        "papers": [],
        "manuscript": None,
        "keyword_sets": [],
        "search_history": [],
        "created_at": now,
        "updated_at": now,
    }
    await db.projects.insert_one(doc)
    doc.pop("_id", None)
    return doc

@projects_router.get("/databases")
async def get_databases(user: dict = Depends(get_current_user)):
    return AVAILABLE_DATABASES

@projects_router.get("/{project_id}")
async def get_project(project_id: str, user: dict = Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]}, {"_id": 0})
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    return project

@projects_router.put("/{project_id}")
async def update_project(project_id: str, update: ProjectUpdate, user: dict = Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    update_data = {k: v for k, v in update.model_dump().items() if v is not None}
    update_data["updated_at"] = datetime.now(timezone.utc).isoformat()
    if "study_summary" in update_data and isinstance(update_data["study_summary"], dict):
        pass  # already a dict from model_dump
    await db.projects.update_one({"id": project_id}, {"$set": update_data})
    return {"message": "Project updated"}

@projects_router.delete("/{project_id}")
async def delete_project(project_id: str, user: dict = Depends(get_current_user)):
    result = await db.projects.delete_one({"id": project_id, "user_id": user["id"]})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Project not found")
    return {"message": "Project deleted"}

# ==================== KEYWORDS ROUTES ====================

@keywords_router.post("/{project_id}")
async def save_keyword_set(project_id: str, keyword_set: KeywordSet, user: dict = Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    keyword_sets = project.get("keyword_sets", [])
    keyword_sets.append(keyword_set.model_dump())
    await db.projects.update_one({"id": project_id}, {"$set": {"keyword_sets": keyword_sets, "updated_at": datetime.now(timezone.utc).isoformat()}})
    return keyword_set

@keywords_router.delete("/{project_id}/{keyword_set_id}")
async def delete_keyword_set(project_id: str, keyword_set_id: str, user: dict = Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    keyword_sets = [ks for ks in project.get("keyword_sets", []) if ks["id"] != keyword_set_id]
    await db.projects.update_one({"id": project_id}, {"$set": {"keyword_sets": keyword_sets}})
    return {"message": "Keyword set deleted"}

# ==================== AGENT ROUTES ====================

async def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text
    except Exception as e:
        logger.error(f"PDF extraction error: {e}")
        return ""

@agents_router.post("/parse-report/{project_id}")
async def parse_report(
    project_id: str,
    file: Optional[UploadFile] = File(None),
    text_content: Optional[str] = Form(None),
    user: dict = Depends(get_current_user)
):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    content = ""
    if file:
        file_bytes = await file.read()
        if file.content_type == "application/pdf" or file.filename.endswith(".pdf"):
            content = await extract_text_from_pdf(file_bytes)
        else:
            content = file_bytes.decode("utf-8", errors="ignore")
    if text_content:
        content += "\n" + text_content

    if not content.strip():
        raise HTTPException(status_code=400, detail="No content provided")

    content = content[:12000]  # Limit to avoid token overflow

    prompt = f"""Extract study information from this clinical/HEOR report. Return ONLY valid JSON.

REPORT:
{content}

Return this JSON structure (use empty string "" if not found):
{{
  "population_setting": "...",
  "sample_size": "...",
  "study_design": "...",
  "inclusion_criteria": "...",
  "exclusion_criteria": "...",
  "interventions": "...",
  "comparators": "...",
  "outcomes": "...",
  "effect_sizes": "...",
  "follow_up": "...",
  "conclusions": "...",
  "tables": [],
  "figures": [],
  "charts": [],
  "raw_tables_text": ""
}}"""

    try:
        response = await call_claude(prompt, "Extract structured data from clinical reports. Return only valid JSON, no markdown.", model="claude-haiku-4-5-20251001")
        json_start = response.find('{')
        json_end = response.rfind('}') + 1
        if json_start == -1:
            raise ValueError("No JSON in response")
        summary_dict = json.loads(response[json_start:json_end])
    except Exception as e:
        logger.error(f"Parse error: {e}")
        summary_dict = {
            "population_setting": "", "sample_size": "", "study_design": "",
            "inclusion_criteria": "", "exclusion_criteria": "", "interventions": "",
            "comparators": "", "outcomes": "", "effect_sizes": "", "follow_up": "",
            "conclusions": content[:500], "tables": [], "figures": [], "charts": [], "raw_tables_text": ""
        }

    # Normalize fields to match StudySummary schema (LLM sometimes returns null or wrong types)
    _str_fields = ["population_setting", "sample_size", "study_design", "inclusion_criteria",
                   "exclusion_criteria", "interventions", "comparators", "outcomes",
                   "effect_sizes", "follow_up", "conclusions", "raw_tables_text"]
    _list_fields = ["tables", "figures", "charts"]
    clean = {}
    for f in _str_fields:
        v = summary_dict.get(f)
        clean[f] = str(v) if v is not None and not isinstance(v, (list, dict)) else (str(v) if v else "")
    for f in _list_fields:
        v = summary_dict.get(f)
        if isinstance(v, list):
            clean[f] = [item if isinstance(item, dict) else {"value": str(item)} for item in v]
        else:
            clean[f] = []
    summary_dict = clean

    await db.projects.update_one(
        {"id": project_id},
        {"$set": {"study_summary": summary_dict, "status": "summary_ready", "updated_at": datetime.now(timezone.utc).isoformat()}}
    )
    return {"study_summary": summary_dict, "message": "Report parsed successfully"}

@agents_router.post("/query-preview/{project_id}")
async def query_preview(project_id: str, filters: SearchFilters, user: dict = Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    dbs = filters.enabled_databases or [k for k, v in AVAILABLE_DATABASES.items() if v["enabled_by_default"]]
    previews = {}
    keywords = " AND ".join(f'"{k}"' for k in filters.main_keywords) if filters.main_keywords else filters.custom_query
    exclusions = " NOT ".join(f'"{k}"' for k in filters.exclusion_keywords) if filters.exclusion_keywords else ""
    mesh = " AND ".join(f'[MeSH Terms]"{m}"' for m in filters.mesh_terms) if filters.mesh_terms else ""

    for db_key in dbs:
        db_info = AVAILABLE_DATABASES.get(db_key, {})
        if db_key in ("pubmed", "pmc"):
            query = keywords
            if mesh:
                query += f" AND {mesh}"
            if exclusions:
                query += f" NOT ({exclusions})"
            if filters.year_from or filters.year_to:
                yr_from = filters.year_from or 2000
                yr_to = filters.year_to or 2025
                query += f" AND {yr_from}:{yr_to}[pdat]"
        else:
            query = keywords
            if exclusions:
                query += f" NOT {exclusions}"

        previews[db_key] = query

    return {"query_previews": previews}

async def search_pubmed(keywords: str, filters: SearchFilters, max_results: int = 20) -> List[Dict]:
    papers = []
    try:
        query = keywords
        if filters.year_from or filters.year_to:
            yr_from = filters.year_from or 2000
            yr_to = filters.year_to or 2025
            query += f" AND {yr_from}:{yr_to}[pdat]"
        if filters.open_access_only:
            query += " AND free full text[sb]"

        async with httpx.AsyncClient(timeout=15) as client:
            search_resp = await client.get(
                "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esearch.fcgi",
                params={"db": "pubmed", "term": query, "retmax": max_results, "retmode": "json", "sort": "relevance"}
            )
            search_data = search_resp.json()
            ids = search_data.get("esearchresult", {}).get("idlist", [])
            if not ids:
                return []

            fetch_resp = await client.get(
                "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi",
                params={"db": "pubmed", "id": ",".join(ids), "retmode": "json"}
            )
            fetch_data = fetch_resp.json()
            result = fetch_data.get("result", {})

            for pmid in ids:
                article = result.get(pmid, {})
                if not article or article.get("error"):
                    continue
                authors_list = article.get("authors", [])
                authors = ", ".join(a.get("name", "") for a in authors_list[:5])
                if len(authors_list) > 5:
                    authors += " et al."

                paper = {
                    "id": str(uuid.uuid4()),
                    "title": article.get("title", "").rstrip("."),
                    "authors": authors,
                    "year": article.get("pubdate", "")[:4],
                    "journal": article.get("source", ""),
                    "pmid": pmid,
                    "doi": next((uid.get("value", "") for uid in article.get("articleids", []) if uid.get("idtype") == "doi"), None),
                    "url": f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/",
                    "abstract": "",
                    "classification": "background",
                    "article_type": "article",
                    "is_open_access": False,
                    "full_text_available": False,
                    "database_source": "pubmed",
                    "databases_found_in": ["pubmed"],
                    "selected": False,
                    "needs_verification": False,
                    "relevance_score": 0.7,
                    "user_notes": "",
                }
                papers.append(paper)
    except Exception as e:
        logger.error(f"PubMed search error: {e}")
    return papers

async def search_semantic_scholar(keywords: str, filters: SearchFilters, max_results: int = 15) -> List[Dict]:
    papers = []
    try:
        params = {"query": keywords, "limit": max_results, "fields": "title,authors,year,venue,externalIds,abstract,isOpenAccess,openAccessPdf"}
        if filters.year_from:
            params["year"] = f"{filters.year_from}-{filters.year_to or 2025}"

        async with httpx.AsyncClient(timeout=15) as client:
            resp = await client.get("https://api.semanticscholar.org/graph/v1/paper/search", params=params)
            data = resp.json()

            for item in data.get("data", []):
                authors_list = item.get("authors", [])
                authors = ", ".join(a.get("name", "") for a in authors_list[:5])
                if len(authors_list) > 5:
                    authors += " et al."

                ext_ids = item.get("externalIds", {}) or {}
                paper = {
                    "id": str(uuid.uuid4()),
                    "title": item.get("title", ""),
                    "authors": authors,
                    "year": str(item.get("year", "")),
                    "journal": item.get("venue", ""),
                    "pmid": ext_ids.get("PubMed"),
                    "doi": ext_ids.get("DOI"),
                    "url": f"https://www.semanticscholar.org/paper/{item.get('paperId', '')}",
                    "abstract": (item.get("abstract") or "")[:500],
                    "classification": "background",
                    "article_type": "article",
                    "is_open_access": item.get("isOpenAccess", False),
                    "full_text_available": bool(item.get("openAccessPdf")),
                    "database_source": "semantic_scholar",
                    "databases_found_in": ["semantic_scholar"],
                    "selected": False,
                    "needs_verification": False,
                    "relevance_score": 0.6,
                    "user_notes": "",
                }
                papers.append(paper)
    except Exception as e:
        logger.error(f"Semantic Scholar error: {e}")
    return papers

async def search_europe_pmc(keywords: str, filters: SearchFilters, max_results: int = 15) -> List[Dict]:
    papers = []
    try:
        query = keywords
        if filters.year_from:
            query += f" AND (FIRST_PDATE:[{filters.year_from}-01-01 TO {filters.year_to or 2025}-12-31])"
        if filters.open_access_only:
            query += " AND OPEN_ACCESS:y"

        async with httpx.AsyncClient(timeout=15) as client:
            resp = await client.get(
                "https://www.ebi.ac.uk/europepmc/webservices/rest/search",
                params={"query": query, "pageSize": max_results, "format": "json", "resultType": "core"}
            )
            data = resp.json()
            for item in data.get("resultList", {}).get("result", []):
                paper = {
                    "id": str(uuid.uuid4()),
                    "title": item.get("title", "").rstrip("."),
                    "authors": item.get("authorString", ""),
                    "year": str(item.get("pubYear", "")),
                    "journal": item.get("journalTitle", ""),
                    "pmid": item.get("pmid"),
                    "doi": item.get("doi"),
                    "url": f"https://europepmc.org/article/{item.get('source', 'MED')}/{item.get('id', '')}",
                    "abstract": (item.get("abstractText") or "")[:500],
                    "classification": "background",
                    "article_type": "article",
                    "is_open_access": item.get("isOpenAccess") == "Y",
                    "full_text_available": item.get("hasPDF") == "Y",
                    "database_source": "europe_pmc",
                    "databases_found_in": ["europe_pmc"],
                    "selected": False,
                    "needs_verification": False,
                    "relevance_score": 0.65,
                    "user_notes": "",
                }
                papers.append(paper)
    except Exception as e:
        logger.error(f"Europe PMC error: {e}")
    return papers

def deduplicate_papers(papers: List[Dict]) -> List[Dict]:
    seen_dois = {}
    seen_pmids = {}
    seen_titles = {}
    result = []

    for paper in papers:
        doi = paper.get("doi")
        pmid = paper.get("pmid")
        title_key = re.sub(r'\W+', '', (paper.get("title") or "").lower())[:50]

        if doi and doi in seen_dois:
            idx = seen_dois[doi]
            result[idx]["databases_found_in"] = list(set(result[idx]["databases_found_in"] + paper["databases_found_in"]))
            continue
        if pmid and pmid in seen_pmids:
            idx = seen_pmids[pmid]
            result[idx]["databases_found_in"] = list(set(result[idx]["databases_found_in"] + paper["databases_found_in"]))
            continue
        if title_key and title_key in seen_titles:
            idx = seen_titles[title_key]
            result[idx]["databases_found_in"] = list(set(result[idx]["databases_found_in"] + paper["databases_found_in"]))
            continue

        idx = len(result)
        result.append(paper)
        if doi:
            seen_dois[doi] = idx
        if pmid:
            seen_pmids[pmid] = idx
        if title_key:
            seen_titles[title_key] = idx

    return result

@agents_router.post("/search-literature/{project_id}")
async def search_literature(project_id: str, filters: SearchFilters, user: dict = Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    keywords = " ".join(filters.main_keywords) if filters.main_keywords else filters.custom_query
    if not keywords.strip():
        raise HTTPException(status_code=400, detail="Please provide keywords or custom query")

    enabled = filters.enabled_databases or [k for k, v in AVAILABLE_DATABASES.items() if v["enabled_by_default"]]
    all_papers = []
    databases_searched = []

    tasks = []
    if "pubmed" in enabled or "pmc" in enabled:
        tasks.append(("pubmed", search_pubmed(keywords, filters, 20)))
        databases_searched.append("pubmed")
    if "semantic_scholar" in enabled:
        tasks.append(("semantic_scholar", search_semantic_scholar(keywords, filters, 15)))
        databases_searched.append("semantic_scholar")
    if "europe_pmc" in enabled:
        tasks.append(("europe_pmc", search_europe_pmc(keywords, filters, 15)))
        databases_searched.append("europe_pmc")

    results = await asyncio.gather(*[t[1] for t in tasks], return_exceptions=True)
    for i, result in enumerate(results):
        if isinstance(result, list):
            all_papers.extend(result)

    all_papers = deduplicate_papers(all_papers)

    # AI classification of papers
    if all_papers and project.get("study_summary"):
        study_summary = project["study_summary"]
        paper_titles = "\n".join(f"{i}. {p['title']}" for i, p in enumerate(all_papers[:20]))
        classify_prompt = f"""Classify these papers relative to the study about: {study_summary.get('interventions', '')} in {study_summary.get('population_setting', '')}.

Papers:
{paper_titles}

Return JSON array:
[{{"index": 0, "classification": "supporting"}}, ...]

Classifications: "supporting" (supports study findings), "contradicting" (contradicts findings), "background" (general background)"""
        try:
            response = await call_claude(classify_prompt, "Classify papers. Return only JSON array.", model="claude-haiku-4-5-20251001")
            json_start = response.find('[')
            json_end = response.rfind(']') + 1
            if json_start != -1:
                classifications = json.loads(response[json_start:json_end])
                for c in classifications:
                    idx = c.get("index")
                    if idx is not None and idx < len(all_papers):
                        all_papers[idx]["classification"] = c.get("classification", "background")
        except Exception as e:
            logger.error(f"Classification error: {e}")

    # Save to project
    await db.projects.update_one(
        {"id": project_id},
        {"$set": {
            "papers": [p if isinstance(p, dict) else p.model_dump() for p in all_papers],
            "status": "search_done",
            "last_search_config": filters.model_dump(),
            "updated_at": datetime.now(timezone.utc).isoformat()
        }}
    )

    return {"papers": all_papers, "count": len(all_papers), "databases_searched": databases_searched}

@agents_router.put("/papers/{project_id}")
async def update_papers(project_id: str, paper_updates: List[Dict], user: dict = Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    papers = project.get("papers", [])
    update_map = {p["id"]: p for p in paper_updates}
    for paper in papers:
        if paper["id"] in update_map:
            u = update_map[paper["id"]]
            paper["selected"] = u.get("selected", paper["selected"])
            paper["classification"] = u.get("classification", paper["classification"])
            paper["user_notes"] = u.get("user_notes", paper.get("user_notes", ""))

    await db.projects.update_one(
        {"id": project_id},
        {"$set": {"papers": papers, "updated_at": datetime.now(timezone.utc).isoformat()}}
    )
    return {"message": "Papers updated"}

# ==================== MANUSCRIPT GENERATION ====================

async def process_manuscript_generation(job_id: str):
    try:
        job = await db.generation_jobs.find_one({"id": job_id}, {"_id": 0})
        if not job:
            return
        await db.generation_jobs.update_one({"id": job_id}, {"$set": {"status": "processing"}})

        project = await db.projects.find_one({"id": job["project_id"]}, {"_id": 0})
        if not project:
            await db.generation_jobs.update_one({"id": job_id}, {"$set": {"status": "failed", "error": "Project not found", "completed_at": datetime.now(timezone.utc).isoformat()}})
            return

        study_summary = project.get("study_summary", {})
        papers = [p for p in project.get("papers", []) if p.get("selected")]

        if not study_summary or not papers:
            await db.generation_jobs.update_one({"id": job_id}, {"$set": {"status": "failed", "error": "Missing study summary or selected papers", "completed_at": datetime.now(timezone.utc).isoformat()}})
            return

        word_counts = job.get("word_counts") or {"abstract": 250, "introduction": 400, "methods": 350, "results": 300, "evidence_comparison": 300, "discussion": 400, "conclusion": 150}

        references = []
        for i, p in enumerate(papers[:15], 1):
            ref = {
                "number": i, "authors": p.get("authors") or "", "title": p.get("title") or "",
                "journal": p.get("journal") or "", "year": p.get("year") or "",
                "volume": p.get("volume") or "", "issue": p.get("issue") or "",
                "pages": p.get("pages") or "", "doi": p.get("doi") or "",
                "pmid": p.get("pmid") or "", "url": p.get("url") or "",
                "abstract": (p.get("abstract") or "")[:200],
                "classification": p.get("classification") or "background",
                "database_source": p.get("database_source") or "",
            }
            references.append(ref)

        refs_text = "\n".join(f"[{r['number']}] {r['authors'][:60]}. \"{r['title'][:100]}\". {r['journal']}. {r['year']}. [{r['classification']}]" for r in references)

        extracted_tables = study_summary.get("tables", [])
        extracted_figures = study_summary.get("figures", [])
        extracted_charts = study_summary.get("charts", [])
        raw_tables_text = study_summary.get("raw_tables_text", "")

        visual_elements_text = ""
        if extracted_tables:
            visual_elements_text += "\n\n## TABLES\n\n"
            for idx, table in enumerate(extracted_tables):
                label = table.get("label", f"Table {idx + 1}")
                caption = table.get("caption", "")
                headers = table.get("headers", [])
                rows = table.get("rows", [])
                visual_elements_text += f"### {label}\n"
                if caption:
                    visual_elements_text += f"**{caption}**\n\n"
                if headers:
                    visual_elements_text += "| " + " | ".join(str(h) for h in headers) + " |\n"
                    visual_elements_text += "| " + " | ".join(["---"] * len(headers)) + " |\n"
                for row in rows:
                    visual_elements_text += "| " + " | ".join(str(c) for c in row) + " |\n"
                visual_elements_text += "\n"

        if extracted_figures:
            visual_elements_text += "\n\n## FIGURES\n\n"
            for idx, figure in enumerate(extracted_figures):
                visual_elements_text += f"### {figure.get('label', f'Figure {idx+1}')}\n"
                if figure.get("caption"):
                    visual_elements_text += f"**{figure['caption']}**\n\n"
                if figure.get("description"):
                    visual_elements_text += f"{figure['description']}\n\n"

        if raw_tables_text:
            visual_elements_text += f"\n\n## ADDITIONAL TABLE DATA\n\n{raw_tables_text}\n"

        manuscript_prompt = f"""Generate a research manuscript as JSON.

STUDY: {study_summary.get('study_design', 'Study')} | N={study_summary.get('sample_size', 'N/A')}
POPULATION: {study_summary.get('population_setting', '')[:150]}
INTERVENTION: {study_summary.get('interventions', '')[:100]}
COMPARATOR: {study_summary.get('comparators', '')[:100]}
OUTCOMES: {study_summary.get('outcomes', '')[:100]}
RESULTS: {study_summary.get('effect_sizes', '')[:150]}
CONCLUSIONS: {study_summary.get('conclusions', '')[:150]}

REFERENCES:
{refs_text}

WORD COUNTS: Abstract ~{word_counts.get('abstract', 250)}, Intro ~{word_counts.get('introduction', 400)}, Methods ~{word_counts.get('methods', 350)}, Results ~{word_counts.get('results', 300)}, Evidence ~{word_counts.get('evidence_comparison', 300)}, Discussion ~{word_counts.get('discussion', 400)}, Conclusion ~{word_counts.get('conclusion', 150)}

RULES:
- Use each reference number in-text like [1], [2] etc.
- Background refs in Introduction, supporting/contradicting in Discussion/Evidence
- Academic tone, no hallucinated citations

Return ONLY this JSON (no markdown, no preamble):
{{"title":"...", "abstract":"...", "introduction":"... [1]...", "methods":"...", "results":"...", "evidence_comparison":"... [2]...", "discussion":"... [3]...", "conclusion":"...", "reference_mapping":{{"introduction":[1], "discussion":[2,3]}}}}"""

        response = await call_claude(manuscript_prompt, "Generate manuscript JSON. Return only valid JSON, no markdown code blocks.")
        json_start = response.find('{')
        json_end = response.rfind('}') + 1
        if json_start == -1 or json_end <= json_start:
            raise ValueError("No JSON in response")
        manuscript_dict = json.loads(response[json_start:json_end])

        manuscript_dict["conclusion"] = manuscript_dict.get("conclusion", "") + "\n\n---\n*AI-generated content — requires human review before publication.*\n---"

        endnote_references = []
        for ref in references:
            endnote_ref = {
                "number": ref["number"], "type": "Journal Article",
                "authors": ref["authors"], "year": ref["year"], "title": ref["title"],
                "journal": ref["journal"], "volume": ref.get("volume", ""),
                "issue": ref.get("issue", ""), "pages": ref.get("pages", ""),
                "doi": ref["doi"], "pmid": ref.get("pmid", ""), "url": ref.get("url", ""),
                "abstract": ref.get("abstract", ""), "classification": ref["classification"],
                "database_source": ref["database_source"],
                "endnote_format": f"{ref['authors']}. {ref['title']}. {ref['journal']}. {ref['year']};{ref.get('volume', '')}" +
                    (f"({ref.get('issue', '')})" if ref.get('issue') else "") +
                    (f":{ref.get('pages', '')}" if ref.get('pages') else "") +
                    (f". doi:{ref['doi']}" if ref['doi'] else "") +
                    (f". PMID:{ref.get('pmid', '')}" if ref.get('pmid') else "")
            }
            endnote_references.append(endnote_ref)

        manuscript = Manuscript(
            title=manuscript_dict.get("title", ""),
            abstract=manuscript_dict.get("abstract", ""),
            introduction=manuscript_dict.get("introduction", ""),
            methods=manuscript_dict.get("methods", ""),
            results=manuscript_dict.get("results", ""),
            evidence_comparison=manuscript_dict.get("evidence_comparison", ""),
            discussion=manuscript_dict.get("discussion", ""),
            conclusion=manuscript_dict.get("conclusion", ""),
            tables_figures=visual_elements_text,
            tables=extracted_tables,
            figures=extracted_figures,
            charts=extracted_charts,
            references=endnote_references,
            reference_mapping=manuscript_dict.get("reference_mapping", {}),
            citation_style=job.get("citation_style", "endnote"),
            version=(project.get("manuscript", {}) or {}).get("version", 0) + 1,
            generated_at=datetime.now(timezone.utc).isoformat()
        )

        await db.projects.update_one(
            {"id": job["project_id"]},
            {"$set": {"manuscript": manuscript.model_dump(), "status": "manuscript_ready", "updated_at": datetime.now(timezone.utc).isoformat()}}
        )
        await db.generation_jobs.update_one(
            {"id": job_id},
            {"$set": {"status": "completed", "result": manuscript.model_dump(), "completed_at": datetime.now(timezone.utc).isoformat()}}
        )
        logger.info(f"Manuscript generation completed for job {job_id}")

    except json.JSONDecodeError as e:
        logger.error(f"JSON error for job {job_id}: {e}")
        await db.generation_jobs.update_one({"id": job_id}, {"$set": {"status": "failed", "error": "Failed to parse AI response", "completed_at": datetime.now(timezone.utc).isoformat()}})
    except Exception as e:
        logger.error(f"Generation error for job {job_id}: {e}")
        await db.generation_jobs.update_one({"id": job_id}, {"$set": {"status": "failed", "error": str(e), "completed_at": datetime.now(timezone.utc).isoformat()}})

@agents_router.post("/generate-manuscript/{project_id}")
async def generate_manuscript(
    project_id: str,
    request: Optional[ManuscriptGenerateRequest] = None,
    sections_to_generate: Optional[str] = None,
    citation_style: str = "endnote",
    user: dict = Depends(get_current_user)
):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]}, {"_id": 0})
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    if not project.get("study_summary"):
        raise HTTPException(status_code=400, detail="Study summary required")
    if not [p for p in project.get("papers", []) if p.get("selected")]:
        raise HTTPException(status_code=400, detail="No papers selected")

    job = GenerationJob(
        project_id=project_id,
        user_id=user["id"],
        citation_style=citation_style,
        sections_to_generate=sections_to_generate.split(',') if sections_to_generate else None,
        word_counts=request.word_counts if request else None,
        status="pending"
    )
    await db.generation_jobs.insert_one(job.model_dump())
    asyncio.create_task(process_manuscript_generation(job.id))
    return {"job_id": job.id, "status": "pending", "message": "Manuscript generation started"}

@agents_router.get("/generation-status/{job_id}")
async def get_generation_status(job_id: str, user: dict = Depends(get_current_user)):
    job = await db.generation_jobs.find_one({"id": job_id, "user_id": user["id"]}, {"_id": 0})
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    response = {"job_id": job["id"], "status": job["status"], "created_at": job["created_at"], "completed_at": job.get("completed_at")}
    if job["status"] == "completed":
        response["manuscript"] = job.get("result")
    elif job["status"] == "failed":
        response["error"] = job.get("error", "Unknown error")
    return response

@agents_router.put("/manuscript/{project_id}")
async def update_manuscript(project_id: str, manuscript_update: Dict, user: dict = Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    manuscript = project.get("manuscript", {}) or {}
    for key in ["title", "abstract", "introduction", "methods", "results", "evidence_comparison", "discussion", "conclusion", "citation_style"]:
        if key in manuscript_update:
            manuscript[key] = manuscript_update[key]
    manuscript["updated_at"] = datetime.now(timezone.utc).isoformat()
    await db.projects.update_one({"id": project_id}, {"$set": {"manuscript": manuscript, "updated_at": datetime.now(timezone.utc).isoformat()}})
    return {"message": "Manuscript updated"}

# ==================== EXPORT ROUTES ====================

@export_router.get("/markdown/{project_id}")
async def export_markdown(project_id: str, user: dict = Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not project or not project.get("manuscript"):
        raise HTTPException(status_code=404, detail="Manuscript not found")
    m = project["manuscript"]
    md = f"# {m.get('title', 'Untitled Manuscript')}\n\n"
    for section, heading in [("abstract", "Abstract"), ("introduction", "Introduction"), ("methods", "Methods"), ("results", "Results"), ("evidence_comparison", "Evidence Comparison"), ("discussion", "Discussion"), ("conclusion", "Conclusion")]:
        md += f"## {heading}\n\n{m.get(section, '')}\n\n"
    md += "## References\n\n"
    for ref in m.get("references", []):
        md += f"[{ref.get('number', '')}] {ref.get('authors', '')}. {ref.get('title', '')}. {ref.get('journal', '')}. {ref.get('year', '')}."
        if ref.get("doi"):
            md += f" DOI: {ref['doi']}"
        md += "\n\n"
    return StreamingResponse(io.BytesIO(md.encode()), media_type="text/markdown", headers={"Content-Disposition": f"attachment; filename=manuscript_{project_id}.md"})

@export_router.get("/docx/{project_id}")
async def export_docx(project_id: str, user: dict = Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not project or not project.get("manuscript"):
        raise HTTPException(status_code=404, detail="Manuscript not found")
    m = project["manuscript"]
    doc = Document()
    title = doc.add_heading(m.get("title", "Untitled Manuscript"), 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%B %d, %Y')}")
    for heading, key in [("Abstract", "abstract"), ("Introduction", "introduction"), ("Methods", "methods"), ("Results", "results"), ("Evidence Comparison", "evidence_comparison"), ("Discussion", "discussion"), ("Conclusion", "conclusion")]:
        doc.add_heading(heading, 1)
        doc.add_paragraph(m.get(key, ""))
    doc.add_heading("References", 1)
    for ref in m.get("references", []):
        ref_text = f"[{ref.get('number', '')}] {ref.get('authors', '')}. {ref.get('title', '')}. {ref.get('journal', '')}. {ref.get('year', '')}."
        if ref.get("doi"):
            ref_text += f" DOI: {ref['doi']}"
        doc.add_paragraph(ref_text)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f"attachment; filename=manuscript_{project_id}.docx"})

@export_router.get("/pdf/{project_id}")
async def export_pdf(project_id: str, user: dict = Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not project or not project.get("manuscript"):
        raise HTTPException(status_code=404, detail="Manuscript not found")
    m = project["manuscript"]
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=72, bottomMargin=72)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='ManTitle', parent=styles['Heading1'], fontSize=18, spaceAfter=30, alignment=TA_CENTER))
    styles.add(ParagraphStyle(name='ManHead', parent=styles['Heading2'], fontSize=14, spaceBefore=20, spaceAfter=10))
    styles.add(ParagraphStyle(name='ManBody', parent=styles['Normal'], fontSize=11, leading=14, alignment=TA_JUSTIFY, spaceAfter=12))
    styles.add(ParagraphStyle(name='ManRef', parent=styles['Normal'], fontSize=10, leading=12, leftIndent=20, firstLineIndent=-20, spaceAfter=6))
    story = [Spacer(1, 60), Paragraph(m.get("title", "Untitled Manuscript"), styles['ManTitle']), Spacer(1, 20), Paragraph(f"Generated: {datetime.now().strftime('%B %d, %Y')}", styles['ManBody']), PageBreak()]
    for heading, key in [("Abstract", "abstract"), ("Introduction", "introduction"), ("Methods", "methods"), ("Results", "results"), ("Evidence Comparison", "evidence_comparison"), ("Discussion", "discussion"), ("Conclusion", "conclusion")]:
        story.append(Paragraph(heading, styles['ManHead']))
        for para in (m.get(key, "") or "").split('\n\n'):
            if para.strip():
                safe = para.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                story.append(Paragraph(safe, styles['ManBody']))
    story.append(PageBreak())
    story.append(Paragraph("References", styles['ManHead']))
    for ref in m.get("references", []):
        ref_text = f"[{ref.get('number', '')}] {ref.get('authors', '')}. {ref.get('title', '')}. {ref.get('journal', '')}. {ref.get('year', '')}."
        if ref.get("doi"):
            ref_text += f" DOI: {ref['doi']}"
        safe = ref_text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
        story.append(Paragraph(safe, styles['ManRef']))
    doc.build(story)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename=manuscript_{project_id}.pdf"})

@export_router.get("/references/{project_id}")
async def export_references(project_id: str, format: str = "json", user: dict = Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not project or not project.get("manuscript"):
        raise HTTPException(status_code=404, detail="Manuscript not found")
    refs = project["manuscript"].get("references", [])
    if format == "endnote":
        content = ""
        for ref in refs:
            content += "%0 Journal Article\n"
            for author in (ref.get("authors", "") or "").split(",")[:5]:
                if author.strip():
                    content += f"%A {author.strip()}\n"
            if ref.get("title"):
                content += f"%T {ref['title']}\n"
            if ref.get("journal"):
                content += f"%J {ref['journal']}\n"
            if ref.get("year"):
                content += f"%D {ref['year']}\n"
            if ref.get("doi"):
                content += f"%R {ref['doi']}\n"
            if ref.get("pmid"):
                content += f"%M {ref['pmid']}\n"
            content += "\n"
        return StreamingResponse(io.BytesIO(content.encode('utf-8')), media_type="text/plain", headers={"Content-Disposition": f"attachment; filename=references_{project_id}.enw"})
    elif format == "csv":
        import csv
        buf = io.StringIO()
        writer = csv.DictWriter(buf, fieldnames=["number", "authors", "title", "journal", "year", "doi", "pmid", "classification", "database_source"], extrasaction='ignore')
        writer.writeheader()
        writer.writerows(refs)
        return StreamingResponse(io.BytesIO(buf.getvalue().encode()), media_type="text/csv", headers={"Content-Disposition": f"attachment; filename=references_{project_id}.csv"})
    else:
        return StreamingResponse(io.BytesIO(json.dumps(refs, indent=2).encode()), media_type="application/json", headers={"Content-Disposition": f"attachment; filename=references_{project_id}.json"})

# ==================== HEALTH ====================

@api_router.get("/")
async def root():
    return {"message": "Manuscript Writer API", "version": "3.0.0"}

@api_router.get("/health")
async def health():
    return {"status": "healthy"}

# Include routers
api_router.include_router(auth_router)
api_router.include_router(projects_router)
api_router.include_router(agents_router)
api_router.include_router(export_router)
api_router.include_router(keywords_router)
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("shutdown")
async def shutdown():
    client.close()
